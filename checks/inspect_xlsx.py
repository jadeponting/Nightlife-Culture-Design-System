#!/usr/bin/env python3
"""Inspect a .xlsx against the system: system faces only, token colours only.

    python3 checks/inspect_xlsx.py file.xlsx

Formula health is recalc.py's job; this checks the presentation layer.
"""
import json, pathlib, sys
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parent.parent
T = json.loads((ROOT / "foundations" / "tokens.json").read_text())
C = T["colour"]
TOKENS = {("FF" + v["hex"][1:]).upper() for v in C["grounds"].values()} \
       | {("FF" + v["hex"][1:]).upper() for v in C["accents"].values() if isinstance(v, dict)} \
       | {("FF" + C["textColours"]["charcoal"]["hex"][1:]).upper(), "00000000"}
FONTS = {"Inter", "Archivo", "Archivo Black", "Archivo ExtraBold", "Archivo SemiBold"}

wb = load_workbook(sys.argv[1])
fails = set()
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            f = c.font
            # A font on an EMPTY cell paints nothing — LibreOffice stamps its
            # Arial default on blank styled cells at every recalc, and no
            # reader can ever see it. Fills paint regardless, so fills are
            # checked on every cell; text properties only where there is text.
            if c.value is not None and f and f.name and f.name not in FONTS:
                fails.add(f"{ws.title}: font “{f.name}” is not a system face")
            if c.value is not None and f and f.color and isinstance(f.color.rgb, str) and f.color.rgb.upper() not in TOKENS:
                fails.add(f"{ws.title}: text colour #{f.color.rgb[2:].upper()} is not a token")
            fill = c.fill
            if fill and fill.fgColor and isinstance(fill.fgColor.rgb, str) \
               and fill.patternType == "solid" and fill.fgColor.rgb.upper() not in TOKENS:
                fails.add(f"{ws.title}: fill #{fill.fgColor.rgb[2:].upper()} is not a token")
print(f"=== {pathlib.Path(sys.argv[1]).name}")
for f in sorted(fails)[:12]:
    print(f"  FAIL  {f}")
if len(fails) > 12: print(f"  … {len(fails)-12} more")
print(f"{len(fails)} failure(s)")
sys.exit(1 if fails else 0)
